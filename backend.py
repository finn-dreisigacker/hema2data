"""
Read blood count PDFs.

Author: Finn Dreisigacker
Date: 28.10.2026

Required packages: pdfplumber and openpyxl

The script checks that the values were read correctly and that there are no parsing errors.
""" 

import argparse
import hashlib
import re
import sys
from datetime import datetime
from pathlib import Path
from pprint import pprint

import pdfplumber

# Settings
EINGABEN = [] # z. B. ["/Users/…/<Name>.pdf"] oder ein Ordner

EXCEL = None # Pfad der Ausgabe-Excel "/Users/…/<excel_name>.xlsx"

EXCEL_BLAETTER = {"Values": False, "Overview": True, "Check": False}

# Setup
PARAMETER = [
    ("WBC", "Leukozyten", "10⁹/L", None),
    ("LYM", "Lymphozyten", "10⁹/L", "%"),
    ("MID", "Mittelgroße Zellen", "10⁹/L", "%"),
    ("GRA", "Granulozyten", "10⁹/L", "%"),
    ("HGB", "Hämoglobin", "g/dL", None),
    ("MCH", "Mittleres korpuskuläres Hämoglobin", "pg", None),
    ("MCHC", "Mittlere korpuskuläre Hämoglobinkonzentration", "g/dL", None),
    ("RBC", "Erythrozyten", "10¹²/L", None),
    ("MCV", "Mittleres korpuskuläres Volumen", "fL", None),
    ("HCT", "Hämatokrit", "%", None),
    ("RDW", "Erythrozytenverteilungsbreite", "%", "fL"),
    ("PLT", "Thrombozyten", "10⁹/L", None),
    ("MPV", "Mittleres Thrombozytenvolumen", "fL", None),
    ("PDW", "Thrombozytenverteilungsbreite", "%", "fL"),
    ("PCT", "Thrombokrit", "%", None),
    ("P-LC", "Anteil großer Thrombozyten", "%", "10⁹/L"),
]

PARAMETER_OHNE_REFERENZ = {"PDW", "PCT", "P-LC"}

EINHEIT_SI = {"10⁹/l": "10⁹/L", "10¹²/l": "10¹²/L", "g/dl": "g/dL", "fl": "fL"}

ID_MUSTER = re.compile(r"RESET-(?:Pilot-)?\d{2}")
                  
WERTE_BREITE = 415

ZAHL = r"-?\d+(?:[.,]\d+)?"
EINHEIT = r"10[⁹¹²]+/[lL]|g/d[lL]|pg|f[lL]|%"

WERTZEILE = re.compile(
    rf"""
    ^
    (?P<code>[A-Z][A-Z-]*)          # Parametercode, z. B. WBC oder P-LC
    (?P<markiert>["*])?             # optionales " oder * (Gerät markiert den Wert)
    \s+
    (?P<wert>{ZAHL})                # gemessener Wert
    \s*
    (?P<einheit>{EINHEIT})          # Einheit des Werts
    (?:                             # optionaler Zusatzwert (z. B. LYM in %)
        \s+
        (?P<zwert>{ZAHL})
        \s*
        (?P<zeinheit>{EINHEIT})
    )?
    \s+
    (?P<min>{ZAHL})                 # Referenz untere Grenze
    \s+
    (?P<max>{ZAHL})                 # Referenz obere Grenze
    \s*
    (?P<refeinheit>{EINHEIT})       # Einheit der Referenz
    $
    """,
    re.VERBOSE,
)

KOPF = re.compile(
    r"""
    ^(Seq\sNo|Date|Profile|Method|Operator|Sample\sID\s1|Sample\sID\s2)
    \b\s*                                                      
    (.*)$                                               
    """,
    re.VERBOSE,
)

PARAM = {code: (einheit, zusatz) for code, _, einheit, zusatz in PARAMETER}

def si(einheit):
    return EINHEIT_SI.get(einheit, einheit) if einheit else einheit

def zahl(s):
    return float(s.replace(",", "."))

## Hilfsfunktionen

# Text aus dem PDF einlesen
def pdf_text(pfad):
    with pdfplumber.open(pfad) as pdf:
        seiten = []
        for s in pdf.pages:
            breite = min(WERTE_BREITE, s.width)
            seiten.append(s.crop((0, 0, breite, s.height)))

        texte = []
        for s in seiten:
            texte.append(s.extract_text() or "")
        return "\n".join(texte), len(pdf.pages)

# Zuordnen des Textes zu Kopf, Werte, Notizen.
def parse_text(text):
    kopf, werte, unerkannt, notiz = {}, [], [], []
    fehler, betroffen, warnungen = [], set(), []
    in_notiz = False
    for zeile in (z.strip() for z in text.splitlines()):
        if not zeile or zeile == "Sample Results":
            continue
        if zeile == "Notes":
            in_notiz = True
        elif in_notiz:
            notiz.append(zeile)
        elif m := KOPF.match(zeile):
            kopf[m[1]] = m[2].strip()
        elif (m := WERTZEILE.match(zeile)) and m["code"] in PARAM:
            werte += werte_aus(m, fehler, betroffen)
            # * ist eine Gerätewarnung; " steht bei PDW/PCT/P-LC auf jedem Befund und bleibt still
            if m["markiert"] == "*":
                warnungen.append(f"{m['code']}: value flagged by the device (*), check the PDF.")
        else:
            unerkannt.append(zeile)

    ## Überprüfung

    # Welche params wurden gefunden?
    codes = []
    for w in werte:
        if " (" not in w["parameter"]:
            codes.append(w["parameter"])

    # Welche params wurden nicht gefunden     
    for c in PARAM:                  
        if c not in codes:
            fehler.append(f"Parameter {c} missing or not recognized.")

    # Wurden doppelte Werte eingelesen? 
    doppelte = set()
    for c in codes:
        if codes.count(c) > 1:
            doppelte.add(c)

    for c in doppelte:
        fehler.append(f"Parameter {c} appears more than once.")
        betroffen.add(c)            

    for z in unerkannt:
        fehler.append(f"Line not recognized: {z}")

    fehler.extend(widersprueche(werte, betroffen))
    
    try:
        datum = datetime.strptime(kopf.get("Date", ""), "%Y-%m-%d %H:%M")
        entnahme = datum.strftime("%Y-%m-%d %H:%M")                 
    except ValueError:           
        entnahme = None
        fehler.append(f"Sample date not readable: {kopf.get('Date')!r}.")
    
    id1 = kopf.get("Sample ID 1", "")
    id2 = kopf.get("Sample ID 2", "")

    teilnehmer = None
    for s in (id1, id2):
        if s:            
            teilnehmer = s
            break
    
    return {
        "collected_at": entnahme,
        "seq_no": kopf.get("Seq No"),
        "sample_id_1": id1,
        "sample_id_2": id2,
        "participant_id": teilnehmer,
        "values": werte,
        "notes": "\n".join(notiz),
        "errors": fehler,
        "error_parameters": sorted(betroffen),
        "warnings": warnungen,
    }

# Werte in jeder Zeile prüfen

def werte_aus(m, fehler, betroffen):
    code = m["code"]
    einheit, zusatz = PARAM[code]

    wert = zahl(m["wert"])
    lo = zahl(m["min"])
    hi = zahl(m["max"])

    vorher = len(fehler)

    if si(m["einheit"]) != einheit or si(m["refeinheit"]) != einheit:
        fehler.append(f"{code}: unit {m['einheit']} instead of {einheit}.")

    if lo >= hi:
        fehler.append(f"{code}: reference range {lo}–{hi} invalid.")

    if wert < 0 or wert > 5 * hi:
        fehler.append(f"{code}: value {wert:g} implausible (reference {lo:g}–{hi:g}).")

    if wert < lo:
        flag = "low"
    elif wert > hi:
        flag = "high"
    else:
        flag = "normal"

    eintrag = {}
    eintrag["parameter"] = code
    eintrag["value"] = wert
    eintrag["unit"] = si(m["einheit"])
    eintrag["ref_min"] = lo
    eintrag["ref_max"] = hi
    eintrag["flag"] = flag
    eintrag["marked"] = bool(m["markiert"])

    if code in PARAMETER_OHNE_REFERENZ:
        eintrag["ref_min"] = None
        eintrag["ref_max"] = None
        eintrag["flag"] = "unknown"

    werte = [eintrag]

    if m["zwert"]:
        if si(m["zeinheit"]) != zusatz:
            fehler.append(f"{code}: secondary value in {si(m['zeinheit'])} instead of {zusatz}.")

        zusatz_eintrag = {}
        zusatz_eintrag["parameter"] = f"{code} ({si(m['zeinheit'])})"
        zusatz_eintrag["value"] = zahl(m["zwert"])
        zusatz_eintrag["unit"] = si(m["zeinheit"])
        zusatz_eintrag["ref_min"] = None
        zusatz_eintrag["ref_max"] = None
        zusatz_eintrag["flag"] = "unknown"
        zusatz_eintrag["marked"] = bool(m["markiert"])
        werte.append(zusatz_eintrag)

    elif zusatz:
        fehler.append(f"{code}: secondary value ({zusatz}) missing.")

    if len(fehler) > vorher:
        betroffen.add(code)

    return werte

# Gegenprüfung: Verrutschte Werte?
def pruefen(fehler, betroffen, name, ist, quelle, soll, codes, rel, abs_):
    toleranz = max(rel * abs(soll), abs_)
    if abs(ist - soll) > toleranz:
        fehler.append(f"{name} = {ist:.4g}, but {quelle} gives {soll:.4g}: "
                      "Value probably misread, check the PDF.")
        for c in codes:
            betroffen.add(c.split(" (")[0])


def widersprueche(werte, betroffen):
    v = {}
    for w in werte:
        v[w["parameter"]] = w["value"]

    fehler = []

    if v.get("RBC") == 0 or v.get("HCT") == 0:
        return fehler

    codes = ("WBC", "LYM", "MID", "GRA")
    if all(c in v for c in codes):
        ist = v["LYM"] + v["MID"] + v["GRA"]
        soll = v["WBC"]
        pruefen(fehler, betroffen, "LYM + MID + GRA", ist, "WBC", soll, codes, 0, 0.25)

    codes = ("LYM (%)", "MID (%)", "GRA (%)")
    if all(c in v for c in codes):
        ist = v["LYM (%)"] + v["MID (%)"] + v["GRA (%)"]
        soll = 100
        pruefen(fehler, betroffen, "LYM % + MID % + GRA %", ist, "100 %", soll, codes, 0, 0.4)

    codes = ("MCH", "HGB", "RBC")
    if all(c in v for c in codes):
        ist = v["MCH"]
        soll = v["HGB"] / v["RBC"] * 10
        pruefen(fehler, betroffen, "MCH", ist, "HGB / RBC", soll, codes, 0.03, 0)

    codes = ("MCHC", "HGB", "HCT")
    if all(c in v for c in codes):
        ist = v["MCHC"]
        soll = v["HGB"] / v["HCT"] * 100
        pruefen(fehler, betroffen, "MCHC", ist, "HGB / HCT", soll, codes, 0.03, 0)

    codes = ("MCV", "HCT", "RBC")
    if all(c in v for c in codes):
        ist = v["MCV"]
        soll = v["HCT"] / v["RBC"] * 10
        pruefen(fehler, betroffen, "MCV", ist, "HCT / RBC", soll, codes, 0.03, 0)

    codes = ("PCT", "PLT", "MPV")
    if all(c in v for c in codes):
        ist = v["PCT"]
        soll = v["PLT"] * v["MPV"] / 10000
        pruefen(fehler, betroffen, "PCT", ist, "PLT × MPV", soll, codes, 0.03, 0.01)

    codes = ("P-LC (10⁹/L)", "PLT", "P-LC")
    if all(c in v for c in codes):
        ist = v["P-LC (10⁹/L)"]
        soll = v["PLT"] * v["P-LC"] / 100
        pruefen(fehler, betroffen, "P-LC (10⁹/L)", ist, "PLT × P-LC %", soll, codes, 0.03, 1)

    return fehler


# Anwenden der Funktionen auf die eingelesene(n) Datei(en)
def lies_pdf(pfad):
    pfad = Path(pfad)
    text, seiten = pdf_text(pfad)
    ergebnis = parse_text(text)
    if seiten != 1:
        ergebnis["warnings"].append(f"PDF has {seiten} pages, expected 1.")
    ergebnis["file"] = pfad.name
    ergebnis["sha256"] = hashlib.sha256(
        pfad.read_bytes()).hexdigest() # Dublettenerkennung
    return ergebnis

# Check (nur beim direkten Start und wenn EINGABEN gefüllt ist; nicht beim Import, z. B. in der Web-App)
if __name__ == "__main__" and EINGABEN:
    ergebnis = lies_pdf(EINGABEN[0])
    pprint(ergebnis, sort_dicts=False)

# Ergebnisse in eine Excel schreiben
def excel_schreiben(ergebnisse, ziel):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    # Steuerzeichen aus Dateiname/PDF-Text würden openpyxl abbrechen lassen
    def text(s):
        return ILLEGAL_CHARACTERS_RE.sub("", s)

    wb = Workbook()

    lang = wb.active
    lang.title = "Values"
    lang.append(["File", "Participant", "Collected at", "Parameter", "Value", "Unit",
                 "Ref min", "Ref max", "Flag", "Error"])

    spalten = []
    for code, _, _, zusatz in PARAMETER:
        spalten.append(code)
        if zusatz:
            spalten.append(f"{code} ({zusatz})")

    breit = wb.create_sheet("Overview")
    breit.append(["File", "Participant", "Collected at", "Status"] + spalten)

    pruef = wb.create_sheet("Check")
    pruef.append(["File", "Type", "Message"])

    for e in ergebnisse:
        datei = text(e["file"])
        proband = text(e["participant_id"] or "")

        if e["errors"]:
            status = "error"
        else:
            status = "ok"

        werte = {}
        for w in e["values"]:
            werte[w["parameter"]] = w

        for w in e["values"]:
            code = w["parameter"].split(" (")[0]
            if code in e["error_parameters"]:
                betroffen = "yes"
            else:
                betroffen = ""
            lang.append([datei, proband, e["collected_at"], w["parameter"], w["value"],
                         w["unit"], w["ref_min"], w["ref_max"], w["flag"], betroffen])

        zeile = [datei, proband, e["collected_at"], status]
        for s in spalten:
            if s in werte:
                zeile.append(werte[s]["value"])
            else:
                zeile.append(None)
        breit.append(zeile)

        for meldung in e["errors"]:
            pruef.append([datei, "error", text(meldung)])
        for meldung in e["warnings"]:
            pruef.append([datei, "warning", text(meldung)])

    for name, aktiv in EXCEL_BLAETTER.items():
        if not aktiv:
            del wb[name]

    # Text aus dem PDF, der mit = beginnt, sonst als Formel speichern (Formel-Injektion)
    for blatt in wb.worksheets:
        for zeile in blatt.iter_rows():
            for zelle in zeile:
                if zelle.data_type == "f":
                    zelle.data_type = "s"
                    zelle.quotePrefix = True  # bleibt Text, auch wenn die Zelle in Excel bearbeitet wird

    wb.save(ziel)


# Ordner und Dateien zu einer Liste von PDFs zusammenführen
def pdfs_sammeln(eingaben):
    pdfs = []
    for eingabe in eingaben:
        eingabe = Path(eingabe)
        if eingabe.is_dir():
            pdfs += sorted(eingabe.glob("*.pdf"))
        else:
            pdfs.append(eingabe)
    return pdfs


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("eingaben", nargs="*", help="PDF files or folders (otherwise EINGABEN above)")
    ap.add_argument("-o", "--excel", help="Save result as Excel file")
    if "IPython" in sys.modules:
        args = ap.parse_args([])
    else:
        args = ap.parse_args()

    if not args.eingaben:
        args.eingaben = EINGABEN
        args.excel = args.excel or EXCEL

    if not args.eingaben:
        print("No PDFs given: set EINGABEN above or pass them as an argument.")
        return 2

    pdfs = pdfs_sammeln(args.eingaben)
    print(f"{len(pdfs)} PDF(s) found.")

    ergebnisse = []
    gesehen = {}

    for pdf in pdfs:
        try:
            e = lies_pdf(pdf)
        except Exception as fehler: # noqa: BLE001
            print(f"[ERROR] {pdf.name}: unreadable ({fehler})")
            ergebnisse.append({"file": pdf.name, "participant_id": None,
                               "collected_at": None, "values": [],
                               "errors": [f"PDF unreadable: {fehler}"],
                               "error_parameters": [], "warnings": []})
            continue

        if e["sha256"] in gesehen:
            e["warnings"].append(f"Same file as {gesehen[e['sha256']]}.")
        else:
            gesehen[e["sha256"]] = e["file"]

        ergebnisse.append(e)

        if e["errors"]:
            zeichen = "[ERROR]"
        else:
            zeichen = "[ OK  ]"

        print(f"{zeichen} {e['file']}: {e['participant_id'] or 'no ID'}, "
              f"{e['collected_at']}, {len(e['values'])} values, {len(e['errors'])} errors")

        for meldung in e["errors"] + e["warnings"]:
            print(f"    {meldung}")

        if not args.excel and len(pdfs) == 1:
            for w in e["values"]:
                print(f"    {w['parameter']:<14}{w['value']:>8g} {w['unit']:<8}{w['flag']}")

    if args.excel and ergebnisse:
        excel_schreiben(ergebnisse, args.excel)
        print(f"Saved: {args.excel}")

    anzahl_fehler = 0
    for e in ergebnisse:
        if e["errors"]:
            anzahl_fehler += 1
    print(f"Done: {len(ergebnisse)} PDF(s), {anzahl_fehler} with errors.")

    if anzahl_fehler > 0:
        return 1
    return 0


if __name__ == "__main__":
    code = main()
    if "IPython" not in sys.modules:
        sys.exit(code)